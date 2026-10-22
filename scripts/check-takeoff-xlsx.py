"""Read browser-exported takeoff workbook independently; requires openpyxl (validation only)."""
import csv
import math
import sys
from openpyxl import load_workbook

book = load_workbook(sys.argv[1], data_only=False, read_only=True)
assert book.sheetnames == ['Takeoff']
sheet = book['Takeoff']
with open(sys.argv[2], encoding='utf-8', newline='') as source:
    expected = list(csv.reader(source))
actual = list(sheet.iter_rows(max_col=17))
assert len(actual) == len(expected) == 3
for index, (cells, row) in enumerate(zip(actual, expected)):
    assert len(cells) == len(row) == 17
    for cell, value in zip(cells, row):
        assert cell.data_type != 'f', 'Unexpected formula cell'
        if cell.data_type == 'n' and cell.value is not None:
            assert math.isclose(cell.value, float(value), rel_tol=1e-14, abs_tol=1e-14)
        else:
            assert (cell.value or '') == value
    if index:
        assert cells[1].data_type == 'n'
        assert cells[5].data_type == 'n'
        assert cells[12].data_type == 'n'
        assert cells[14].data_type == 'n'
assert not book._external_links
book.close()
print('Independent openpyxl reader: 2 quantities, numeric cells, evidence and no formulas verified against CSV.')
